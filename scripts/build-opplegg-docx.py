#!/usr/bin/env python3
"""Bygg nedlastbar Word-fil for undervisningsopplegget 2A og 2B."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "public" / "fagstoff" / "undervisningsopplegg-2a-2b.docx"
SITE = "https://matteguiden.no"
GREEN = RGBColor(0x0F, 0x76, 0x6E)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
LINK = "2563EB"


def shade(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, *, size=11, bold=False, color=INK, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:cs"), "Calibri")


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def para(doc, text="", *, size=11, bold=False, color=INK, space_after=8, space_before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = GREEN
        run.font.name = "Calibri"


def bullet(doc, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text)
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)


def task_link(doc, label: str, path: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    add_hyperlink(p, label, f"{SITE}{path}")
    r = p.add_run(f"  —  {SITE}{path}")
    set_run_font(r, size=10, color=MUTED)


def video_row(table, title: str, meta: str, url: str) -> None:
    row = table.add_row().cells
    row[0].text = ""
    row[1].text = ""
    row[2].text = ""
    p0 = row[0].paragraphs[0]
    r = p0.add_run(title)
    set_run_font(r, bold=True, size=11)
    p1 = row[1].paragraphs[0]
    r = p1.add_run(meta)
    set_run_font(r, size=10, color=MUTED)
    p2 = row[2].paragraphs[0]
    add_hyperlink(p2, url, url)


def box(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, "ECFDF5")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + " ")
    set_run_font(r, bold=True, size=11, color=GREEN)
    r = p.add_run(body)
    set_run_font(r, size=11)
    para(doc, "", space_after=10)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Matteguiden 1T  ·  Undervisningsopplegg 2A og 2B")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = footer.add_run("Oppgavene ligger i Matteguiden: ")
    set_run_font(r, size=9, color=MUTED)
    add_hyperlink(footer, f"{SITE}/oppgaver/kap2/2A", f"{SITE}/oppgaver/kap2/2A")
    r = footer.add_run("  og  ")
    set_run_font(r, size=9, color=MUTED)
    add_hyperlink(footer, f"{SITE}/oppgaver/kap2/2B", f"{SITE}/oppgaver/kap2/2B")

    title = para(doc, "Undervisningsopplegg", size=28, bold=True, color=GREEN, space_after=2)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(
        doc,
        "2A Forenkle uttrykk  ·  2B Kvadratsetninger",
        size=16,
        bold=True,
        color=INK,
        space_after=4,
    )
    para(
        doc,
        "Fire økter i matematikk 1T, bygd på oppgavetypene i Matteguiden. Se video først, øv deretter.",
        size=12,
        color=MUTED,
        space_after=14,
    )

    heading(doc, "Ramme for opplegget", 1)
    bullet(doc, " Matematikk 1T, kapittel 2 Algebra og mønstre", bold_prefix="Fag og trinn:")
    bullet(doc, " 4 økter à 60–90 minutter (kan kjøres som to dobbelttimer)", bold_prefix="Varighet:")
    bullet(
        doc,
        " 1A regnerekkefølge, 1C brøk og 1D potenser. Elevene må kunne x · x = x².",
        bold_prefix="Forkunnskap:",
    )
    bullet(
        doc,
        " 2C felles faktor og 2D faktorisering (kvadratsetning baklengs)",
        bold_prefix="Neste steg:",
    )
    bullet(
        doc,
        " rekne med bokstavuttrykk, parentesuttrykk og kvadratiske uttrykk, og bruke kvadratsetningane",
        bold_prefix="Kompetansemål (LK20, 1T):",
    )
    box(
        doc,
        "Arbeidsform i hver økt:",
        "8–12 min video → 5 min tavleeksempel (samme type som oppgavene) → øving i Matteguiden med hint → 5 min «finn feilen» → avslutning med 3–4 quizspørsmål.",
    )

    heading(doc, "Økt 1 — Like ledd og distributiv lov", 1)
    para(
        doc,
        "Læringsmål. Eleven kan samle like ledd, gange et tall eller en bokstav inn i en parentes, og bytte fortegn når det står minus foran parentesen.",
        space_after=8,
    )
    para(
        doc,
        "Dette er oppgavetypen i 2.1, 2.17 og starten av 2.2: 7x − 2 + 3x + 8, 4(−x + 4) og −2(3x − 4).",
        space_after=8,
    )
    heading(doc, "Videoer før øving", 2)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Vi trekker sammen ledd med bokstaver", "https://www.youtube.com/watch?v=PA7Q18hLK0E")
    r = p.add_run("  —  MatematikkMedBirger")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=8)
    add_hyperlink(p, "Sinus 1T kap 1.1 Regnerekkefølge", "https://www.youtube.com/watch?v=HCvi7QZBoGE")
    r = p.add_run("  —  Johan Hake (parenteser og rekkefølge)")
    set_run_font(r, size=11, color=MUTED)
    heading(doc, "Tavleeksempel (samme type som 2.1 og 2.17)", 2)
    para(doc, "3(x − 4) − (x + 2) = 3x − 12 − x − 2 = 2x − 14", bold=True, space_after=6)
    para(
        doc,
        "Si høyt: «Tallet utenfor ganges med alle ledd. Minus foran parentes er det samme som å gange med −1.»",
        space_after=8,
    )
    heading(doc, "Oppgaver i Matteguiden", 2)
    task_link(doc, "2.1 — Trekk sammen like ledd", "/oppgaver/kap2/2A#oppgave-2.1")
    task_link(doc, "2.17 — Like ledd og minus foran parentes", "/oppgaver/kap2/2A#oppgave-2.17")
    task_link(doc, "2.2 a–b — Løs opp parentesene", "/oppgaver/kap2/2A#oppgave-2.2")
    task_link(doc, "2.18 a–b — Parenteser og fortegn", "/oppgaver/kap2/2A#oppgave-2.18")
    bullet(doc, "Ekstra: 2.38 a–d (samme øving, flere ledd)")
    box(
        doc,
        "Vanlig feil:",
        "2(x − 4) = 2x − 4 (glemmer å gange −4). Rett: 2x − 8. Sjekk med x = 5: venstre side 2 · 1 = 2, høyre 10 − 4 = 6.",
    )

    heading(doc, "Økt 2 — To parenteser, areal og «finn feilen»", 1)
    para(
        doc,
        "Læringsmål. Eleven kan gange to parenteser (a + b)(c + d), trekke sammen, og lage et forenklet uttrykk for omkrets og areal. Eleven kan avsløre en typisk parentesfeil ved å sette inn et tall.",
        space_after=8,
    )
    para(
        doc,
        "Dette er 2.19, 2.3, 2.18 c og 2.20: hvert ledd ganges med hvert ledd, konjugat som dukker opp tidlig, ramme/rektangel, og feiljakt.",
        space_after=8,
    )
    heading(doc, "Videoer før øving", 2)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Første kvadratsetning", "https://www.youtube.com/watch?v=LgDQ29AAfEM")
    r = p.add_run("  —  Nils Mulder (viser (a + b)(a + b) ledd for ledd)")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=8)
    add_hyperlink(p, "Å lage formler og uttrykk", "https://www.youtube.com/watch?v=scwkoWmWDBY")
    r = p.add_run("  —  Stian Ripe Daviknes (tekst til formel)")
    set_run_font(r, size=11, color=MUTED)
    heading(doc, "Tavleeksempel (samme type som 2.19)", 2)
    para(doc, "(x + 2)(x + 3) = x² + 3x + 2x + 6 = x² + 5x + 6", bold=True, space_after=6)
    para(
        doc,
        "Tegn et rektangel med sider x + 2 og x + 3. De fire feltene er x², 3x, 2x og 6 — samme fire ledd som i utregningen.",
        space_after=8,
    )
    heading(doc, "Oppgaver i Matteguiden", 2)
    task_link(doc, "2.19 — To parenteser", "/oppgaver/kap2/2A#oppgave-2.19")
    task_link(doc, "2.3 — To parenteser og konjugat", "/oppgaver/kap2/2A#oppgave-2.3")
    task_link(doc, "2.18 c — Omkrets av rektangel", "/oppgaver/kap2/2A#oppgave-2.18")
    task_link(doc, "2.20 — Finn feilen og sett inn tall", "/oppgaver/kap2/2A#oppgave-2.20")
    bullet(doc, "Ekstra: 2.39 (omkrets, areal og sjekk av elevsvar) og 2.40 (feiljakt)")
    box(
        doc,
        "Vanlig feil:",
        "(x + 2)(x + 3) = x² + 6 (glemmer mellomleddene). Rett: x² + 5x + 6. Quiz i 2A avrunder økta.",
    )

    heading(doc, "Økt 3 — De tre kvadratsetningene", 1)
    para(
        doc,
        "Læringsmål. Eleven kan skrive ned og bruke de tre setningene, og kvadrere hele det første leddet, for eksempel (3a)² = 9a².",
        space_after=8,
    )
    para(doc, "(a + b)² = a² + 2ab + b²", bold=True, space_after=2)
    para(doc, "(a − b)² = a² − 2ab + b²", bold=True, space_after=2)
    para(doc, "(a + b)(a − b) = a² − b²", bold=True, space_after=8)
    para(doc, "Dette er 2.4, 2.5 og 2.21.", space_after=8)
    heading(doc, "Videoer før øving (én setning om gangen)", 2)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Første kvadratsetning", "https://www.youtube.com/watch?v=LgDQ29AAfEM")
    r = p.add_run("  —  Nils Mulder")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Andre kvadratsetning", "https://www.youtube.com/watch?v=vKmPyzXwm_A")
    r = p.add_run("  —  Nils Mulder")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=8)
    add_hyperlink(p, "Konjugatsetningen", "https://www.youtube.com/watch?v=isqqIJYcnB0")
    r = p.add_run("  —  Nils Mulder (tredje kvadratsetning)")
    set_run_font(r, size=11, color=MUTED)
    heading(doc, "Tavleeksempel (samme type som 2.5)", 2)
    para(doc, "(3a − 2)² = 9a² − 12a + 4", bold=True, space_after=2)
    para(doc, "(4y + 3)(4y − 3) = 16y² − 9", bold=True, space_after=6)
    para(
        doc,
        "La elevene si a og b høyt før de setter inn. Motbevis myten (a + b)² = a² + b² med a = 2, b = 3: 25 ≠ 13. Mellomleddet 2ab mangler.",
        space_after=8,
    )
    heading(doc, "Oppgaver i Matteguiden", 2)
    task_link(doc, "2.4 — Skriv ned kvadratsetningene", "/oppgaver/kap2/2B#oppgave-2.4")
    task_link(doc, "2.5 — Bruk setningene", "/oppgaver/kap2/2B#oppgave-2.5")
    task_link(doc, "2.21 — Første, andre og konjugat", "/oppgaver/kap2/2B#oppgave-2.21")
    bullet(doc, "Geometri (hvorfor 2ab): ekstra 2.42 — kvadrat med side a + b")
    box(
        doc,
        "Vanlig feil:",
        "(2a + 1)² = 2a² + 1 eller 4a² + 1. Rett: 4a² + 4a + 1. Både 2a og 1 skal kvadreres, og dobbeltproduktet er 4a.",
    )

    heading(doc, "Økt 4 — Hoderegning, røtter og sammensatte uttrykk", 1)
    para(
        doc,
        "Læringsmål. Eleven kjenner igjen setningene i nye forkledninger: 103² = (100 + 3)², 98 · 102 = (100 − 2)(100 + 2), (√7 − 1)(√7 + 1) og (x + y)² − (x − y)².",
        space_after=8,
    )
    para(doc, "Dette er 2.22, 2.6, 2.23 og 2.24.", space_after=8)
    heading(doc, "Videoer før øving", 2)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Konjugatsetningen", "https://www.youtube.com/watch?v=isqqIJYcnB0")
    r = p.add_run("  —  Nils Mulder (bruk til a² − b²)")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=2)
    add_hyperlink(p, "Regneregler: Kvadratrot", "https://www.youtube.com/watch?v=NRx60-H6ZY0")
    r = p.add_run("  —  Hege Rødsdalen  ((√a)² = a)")
    set_run_font(r, size=11, color=MUTED)
    p = para(doc, "", space_after=8)
    add_hyperlink(p, "Gjenkjenne fullstendige kvadrater", "https://www.youtube.com/watch?v=B9vsG1QMy1o")
    r = p.add_run("  —  Nils Mulder (bro mot 2D)")
    set_run_font(r, size=11, color=MUTED)
    heading(doc, "Tavleeksempel (samme type som 2.22 og 2.24)", 2)
    para(doc, "103² = (100 + 3)² = 10 000 + 600 + 9 = 10 609", bold=True, space_after=2)
    para(doc, "(√7 − 1)(√7 + 1) = 7 − 1 = 6", bold=True, space_after=8)
    heading(doc, "Oppgaver i Matteguiden", 2)
    task_link(doc, "2.22 — Hoderegning med kvadratsetninger", "/oppgaver/kap2/2B#oppgave-2.22")
    task_link(doc, "2.6 — Røtter og sammensatte uttrykk", "/oppgaver/kap2/2B#oppgave-2.6")
    task_link(doc, "2.23 — Forenkle med setningene", "/oppgaver/kap2/2B#oppgave-2.23")
    task_link(doc, "2.24 — Kvadratsetninger med rot", "/oppgaver/kap2/2B#oppgave-2.24")
    bullet(doc, "Ekstra: 2.43 b (103 · 97) og 2.41 e")
    para(doc, "Avslutt med quiz i 2B. Elever som er ferdige, kan starte 2C Felles faktor.", space_after=8)
    box(
        doc,
        "Sjekk ut av 2A + 2B:",
        "Eleven forenkler 3(x − 2) + x, ganger ut (x + 2)(x + 3), skriver (x + 4)² uten å glemme 8x, og regner 21 · 19 som 20² − 1.",
    )

    heading(doc, "Alle YouTube-lenker i opplegget", 1)
    para(
        doc,
        "Titlene er sjekket mot selve YouTube-videoene. Trykk på lenken for å åpne klippet.",
        space_after=8,
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, label in enumerate(("Video", "Kanal", "Lenke")):
        hdr[i].text = ""
        shade(hdr[i], "0F766E")
        p = hdr[i].paragraphs[0]
        r = p.add_run(label)
        set_run_font(r, bold=True, size=11, color=RGBColor(0xFF, 0xFF, 0xFF))
    videos = [
        ("Vi trekker sammen ledd med bokstaver", "MatematikkMedBirger", "https://www.youtube.com/watch?v=PA7Q18hLK0E"),
        ("Sinus 1T kap 1.1 Regnerekkefølge", "Johan Hake", "https://www.youtube.com/watch?v=HCvi7QZBoGE"),
        ("Å lage formler og uttrykk", "Stian Ripe Daviknes", "https://www.youtube.com/watch?v=scwkoWmWDBY"),
        ("Første kvadratsetning", "Nils Mulder", "https://www.youtube.com/watch?v=LgDQ29AAfEM"),
        ("Andre kvadratsetning", "Nils Mulder", "https://www.youtube.com/watch?v=vKmPyzXwm_A"),
        ("Konjugatsetningen", "Nils Mulder", "https://www.youtube.com/watch?v=isqqIJYcnB0"),
        ("Regneregler: Kvadratrot", "Hege Rødsdalen", "https://www.youtube.com/watch?v=NRx60-H6ZY0"),
        ("Gjenkjenne fullstendige kvadrater", "Nils Mulder", "https://www.youtube.com/watch?v=B9vsG1QMy1o"),
    ]
    for title, meta, url in videos:
        video_row(table, title, meta, url)

    widths = (Cm(7.2), Cm(4.2), Cm(6.4))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    para(doc, "", space_after=6)
    p = para(doc, "Opplegget i Matteguiden: ", size=10, color=MUTED, space_before=12)
    add_hyperlink(p, f"{SITE}/ressurser?t=opplegg-2a-2b", f"{SITE}/ressurser?t=opplegg-2a-2b")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Skrev {path} ({path.stat().st_size} byte)")
